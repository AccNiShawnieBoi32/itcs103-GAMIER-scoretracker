import openpyxl
from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "student_score.xlsx"

def initialize_workbook():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Scores"
        ws.append(["Name", "Score", "Status"])
        wb.save(FILE_NAME)

def get_status(score):
    return "Pass" if score >= 75 else "Fail"

def add_or_update_student(name, score):
    initialize_workbook()
    wb = load_workbook(FILE_NAME)
    ws = wb.active

    updated = False
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value == name:
            row[1].value = score
            row[2].value = get_status(score)
            updated = True
            break

    if not updated:
        ws.append([name, score, get_status(score)])

    wb.save(FILE_NAME)
    print(f"Record {'updated' if updated else 'added'} for {name}.")

def display_all_records():
    initialize_workbook()
    wb = load_workbook(FILE_NAME)
    ws = wb.active

    print(f"\n{'Name':<20}{'Score':<10}{'Status'}")
    print("-" * 40)
    for row in ws.iter_rows(min_row=2, values_only=True):
        print(f"{row[0]:<20}{row[1]:<10}{row[2]}")

def main():
    while True:
        print("\n1. Add/Update Student Score")
        print("2. Display All Records")
        print("3. Exit")

        choice = input("Enter your choice (1-3): ").strip()

        if choice == '1':
            name = input("Enter student name: ").strip()
            try:
                score = int(input("Enter student score (0-100): ").strip())
                if 0 <= score <= 100:
                    add_or_update_student(name, score)
                else:
                    print("Score must be between 0 and 100.")
            except ValueError:
                print("Invalid input. Please enter a numeric score.")
        elif choice == '2':
            display_all_records()
        elif choice == '3':
            print("Exiting program.")
            break
        else:
            print("Invalid choice. Please enter 1, 2, or 3.")

if __name__ == "__main__":
    main()
